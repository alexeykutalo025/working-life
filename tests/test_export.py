"""Exports, and the integrity statement none of them may leave without.

Spec 9.6: "An export that silently omits known problems is a defect, not a
nicety." These tests treat it as one.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

import pytest

from recall.config import Settings
from recall.export import CALENDAR_COLUMNS, ExportError, run_export
from recall.export.base import ExportSelection, build_statement
from recall.export.csv_export import _safe
from recall.extract import Extractor, parse_kinds
from recall.integrity.engine import Finding, record_finding
from recall.scan.walker import Scanner
from tests.fixtures.generate import generate_ics, generate_vcs


@pytest.fixture
def archive(tmp_path: Path, settings: Settings, conn):
    """A small archive of calendar entries, read from real fixture files."""
    fixtures = tmp_path / "cal"
    generate_ics(fixtures)
    generate_vcs(fixtures)
    Scanner(settings, conn).run([fixtures])
    Extractor(settings, conn).run(kinds=parse_kinds("calendar"))
    return conn


# --- the statement is not optional ---------------------------------------


def test_csv_always_gets_a_statement_beside_it(archive, settings):
    data, statement_path, _ = run_export(archive, settings, fmt="csv", kind="calendar")
    assert data.exists()
    assert statement_path is not None and statement_path.exists()
    assert statement_path.name.endswith("_integrity.txt")


def test_markdown_always_gets_a_statement_beside_it(archive, settings):
    _, statement_path, _ = run_export(archive, settings, fmt="markdown", kind="calendar")
    assert statement_path is not None and statement_path.exists()


def test_xlsx_carries_the_statement_as_its_first_sheet(archive, settings):
    openpyxl = pytest.importorskip("openpyxl")
    data, statement_path, _ = run_export(archive, settings, fmt="xlsx", kind="calendar")
    assert statement_path is None, "the workbook carries it inside"
    workbook = openpyxl.load_workbook(data)
    assert workbook.sheetnames[0] == "Integrity"
    assert "Records" in workbook.sheetnames


def test_json_carries_a_findings_key(archive, settings):
    data, statement_path, _ = run_export(archive, settings, fmt="json", kind="calendar")
    assert statement_path is None
    document = json.loads(data.read_text(encoding="utf-8"))
    assert "findings" in document, "spec 9.6 names this key exactly"
    assert "records" in document
    assert document["record_count"] == len(document["records"])


def test_the_statement_names_what_is_missing(archive, settings):
    source_id = archive.execute("SELECT id FROM source_files LIMIT 1").fetchone()["id"]
    record_finding(archive, Finding(
        code="partial_parse", severity="critical",
        title="archive1998.pst: about 8,000 records could not be read",
        detail="This file reports 12,400 messages. We could read 4,380.",
        source_file_id=source_id, estimated_loss=8000,
    ))

    _, statement_path, statement = run_export(archive, settings, fmt="csv", kind="calendar")
    text = statement_path.read_text(encoding="utf-8")

    assert statement.estimated_missing == 8000
    assert "8,000" in text
    assert "COULD NOT BE READ" in text.upper()
    assert "archive1998.pst" in text


def test_a_clean_export_says_so_rather_than_saying_nothing(archive, settings):
    """Silence would read as "nobody checked"."""
    archive.execute("DELETE FROM findings")
    _, statement_path, statement = run_export(archive, settings, fmt="csv", kind="calendar")
    text = statement_path.read_text(encoding="utf-8")
    assert statement.is_clean
    assert "NO KNOWN PROBLEMS" in text
    assert "cannot know about a mailbox that was never on this computer" in text


def test_the_statement_counts_what_was_actually_written(archive, settings):
    data, statement_path, statement = run_export(archive, settings, fmt="csv", kind="calendar")
    with open(data, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert len(rows) == statement.exported_count


def test_undated_records_are_named_in_the_statement(archive, settings):
    archive.execute(
        "INSERT INTO items(kind, dedup_key, occurred_utc, subject) "
        "VALUES ('event', 'undated-1', NULL, 'No date on this one')"
    )
    _, statement_path, statement = run_export(archive, settings, fmt="csv", kind="calendar")
    text = statement_path.read_text(encoding="utf-8")
    assert statement.undated_count >= 1
    assert "no usable date" in text
    assert "no date has been invented" in text


# --- the calendar CSV, which must open in Excel unedited ------------------


def test_the_spec_columns_come_first_and_in_order(archive, settings):
    data, _, _ = run_export(archive, settings, fmt="csv", kind="calendar")
    with open(data, encoding="utf-8-sig", newline="") as fh:
        header = next(csv.reader(fh))
    assert header[: len(CALENDAR_COLUMNS)] == CALENDAR_COLUMNS


def test_basic_export_writes_only_the_spec_columns(archive, settings):
    data, _, _ = run_export(archive, settings, fmt="csv", kind="calendar", full=False)
    with open(data, encoding="utf-8-sig", newline="") as fh:
        header = next(csv.reader(fh))
    assert header == CALENDAR_COLUMNS


def test_csv_has_a_bom_so_excel_reads_utf8(archive, settings):
    """Without it, Excel on Windows turns every accented name into mojibake."""
    data, _, _ = run_export(archive, settings, fmt="csv", kind="calendar")
    assert data.read_bytes().startswith(b"\xef\xbb\xbf")


def test_accented_text_survives_the_round_trip(archive, settings):
    data, _, _ = run_export(archive, settings, fmt="csv", kind="calendar")
    text = data.read_text(encoding="utf-8-sig")
    assert "Réunion avec Aoife" in text
    assert "Église" in text


def test_dates_and_times_are_separate_and_iso(archive, settings):
    data, _, _ = run_export(archive, settings, fmt="csv", kind="calendar")
    with open(data, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    dated = [r for r in rows if r["date"]]
    assert dated
    for row in dated:
        assert len(row["date"]) == 10 and row["date"][4] == "-"
        if row["start"]:
            assert len(row["start"]) == 5 and row["start"][2] == ":"


def test_duration_is_a_plain_number_of_minutes(archive, settings):
    data, _, _ = run_export(archive, settings, fmt="csv", kind="calendar")
    with open(data, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    durations = [r["duration_min"] for r in rows if r["duration_min"]]
    assert durations
    assert all(d.isdigit() for d in durations), "Excel must be able to sum this column"
    assert "90" in durations


def test_an_unknown_duration_is_blank_not_zero(archive, settings):
    """A duration we do not know is not a duration of zero."""
    archive.execute(
        "INSERT INTO items(kind, dedup_key, occurred_utc, end_utc, subject) "
        "VALUES ('event', 'noend', '2005-01-01T10:00:00Z', NULL, 'No end time')"
    )
    data, _, _ = run_export(archive, settings, fmt="csv", kind="calendar")
    with open(data, encoding="utf-8-sig", newline="") as fh:
        row = next(r for r in csv.DictReader(fh) if r["subject"] == "No end time")
    assert row["duration_min"] == ""


def test_an_unknown_timezone_is_said_so_in_the_row(archive, settings):
    """The honest-count rule at the level of one row."""
    data, _, _ = run_export(archive, settings, fmt="csv", kind="calendar")
    with open(data, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    floating = [r for r in rows if r["timezone"] == "unknown"]
    assert floating, "the fixtures contain floating times"
    for row in floating:
        assert row["timezone_known"] == "no"
        assert "timezone not recorded" in row["data_quality"]


def test_attendees_are_one_cell_with_a_count_beside_it(archive, settings):
    data, _, _ = run_export(archive, settings, fmt="csv", kind="calendar")
    with open(data, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    meeting = next(r for r in rows if r["subject"].startswith("Review Q1"))
    assert meeting["attendee_count"] == "2"
    assert meeting["attendees"].count(";") == 1
    assert "mobrien@contractmktg.com" in meeting["attendees"]


def test_the_source_file_is_named_on_every_row(archive, settings):
    data, _, _ = run_export(archive, settings, fmt="csv", kind="calendar")
    with open(data, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert all(r["source_file"] for r in rows)


# --- formula injection ----------------------------------------------------


@pytest.mark.parametrize("dangerous", [
    "=1+1",
    "=cmd|'/c calc'!A1",
    "+44 20 7946 0000",
    "-- Tim McCarthy",
    "@SUM(A1:A9)",
])
def test_formula_starters_are_neutralised(dangerous):
    """Excel executes a cell starting with these. Real mail contains them."""
    assert _safe(dangerous).startswith("'")
    assert dangerous in _safe(dangerous)


def test_ordinary_text_is_not_prefixed():
    assert _safe("Quarterly figures") == "Quarterly figures"
    assert _safe("2003-04-14") == "2003-04-14"


def test_none_becomes_an_empty_cell_not_the_word_none():
    assert _safe(None) == ""


def test_booleans_read_as_words():
    assert _safe(True) == "yes"
    assert _safe(False) == "no"


# --- refusals -------------------------------------------------------------


def test_an_unknown_format_is_refused_by_name(archive, settings):
    with pytest.raises(ExportError, match="not a format"):
        run_export(archive, settings, fmt="pdf", kind="calendar")


def test_an_unknown_kind_is_refused_by_name(archive, settings):
    with pytest.raises(ExportError, match="not something Recall can export"):
        run_export(archive, settings, fmt="csv", kind="recipes")


def test_writing_into_onedrive_is_refused(archive, settings, tmp_path, monkeypatch):
    """Spec section 14: never write into a OneDrive folder."""
    from recall.scan.onedrive import HydrationRefused

    onedrive = tmp_path / "OneDrive"
    onedrive.mkdir()
    monkeypatch.setenv("OneDrive", str(onedrive))
    with pytest.raises(HydrationRefused, match="OneDrive"):
        run_export(archive, settings, fmt="csv", kind="calendar",
                   out_path=onedrive / "calendar")


# --- the statement itself -------------------------------------------------


def test_statement_dict_is_serialisable(conn, settings):
    statement = build_statement(conn, ExportSelection(), 0, settings.db_path)
    assert json.dumps(statement.as_dict())


def test_statement_rows_are_label_value_pairs(conn, settings):
    statement = build_statement(conn, ExportSelection(), 42, settings.db_path)
    rows = statement.rows()
    assert ("Records exported", "42") in rows
    assert any("complete" in label.lower() for label, _ in rows)


# --- Excel's own rules, which real mail breaks ----------------------------


def test_xlsx_survives_control_characters_in_real_mail(archive, settings):
    """openpyxl refuses to write them and Excel will not open a file with them.

    Real mail is full of stray control bytes from thirty years of mail clients,
    so an export that crashes on them is an export that does not work. This was
    found by running the export against the real .ost on the build machine.
    """
    pytest.importorskip("openpyxl")

    bell = chr(7)
    null = chr(0)
    unit_separator = chr(31)

    archive.execute(
        "INSERT INTO items(kind, dedup_key, subject, body_text, occurred_utc) "
        "VALUES ('event', 'control-chars', ?, ?, '2005-06-01T09:00:00Z')",
        (
            f"Subject with {bell} a bell",
            f"Body with {null} a null and {unit_separator} a unit separator",
        ),
    )

    data, _, _ = run_export(archive, settings, fmt="xlsx", kind="calendar")
    assert data.exists()

    import openpyxl

    workbook = openpyxl.load_workbook(data)
    assert "Records" in workbook.sheetnames

    subjects = [
        cell.value
        for row in workbook["Records"].iter_rows(min_row=2)
        for cell in row
        if isinstance(cell.value, str) and "a bell" in cell.value
    ]
    assert subjects, "the record is still there"
    assert bell not in subjects[0], "the control character is gone"


def test_xlsx_neutralises_formula_injection():
    """Excel treats a leading = as a formula in a workbook just as in a CSV."""
    pytest.importorskip("openpyxl")
    from recall.export.xlsx_export import _cell

    assert _cell("=1+1").startswith("'")
    assert _cell("-- Tim McCarthy").startswith("'")
    assert _cell("Quarterly figures") == "Quarterly figures"


def test_xlsx_truncates_a_cell_too_long_for_excel():
    """A cell cannot hold more than 32,767 characters, and a mail body can."""
    pytest.importorskip("openpyxl")
    from recall.export.xlsx_export import _cell

    result = _cell("x" * 40_000)
    assert len(result) < 33_000
    assert "truncated for Excel" in result
